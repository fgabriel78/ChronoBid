"""Generador de plantilla y lector/escritor resiliente de hojas Excel (.xlsx)."""

from __future__ import annotations

import os
import tempfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from chronobid.models.auction import (
    ERROR_TITLE_PHRASES,
    AuctionStatus,
    GlobalBudgetState,
    ReserveStatus,
    WatchLot,
    extract_lot_id,
    is_valid_catawiki_url,
)

HEADER_ROW = 6
FIRST_DATA_ROW = 7

TABLE_HEADERS: dict[str, str] = {
    "A": "URL Catawiki",
    "B": "Puja Máxima (EUR)",
    "C": "Título del Reloj",
    "D": "Hora de Cierre",
    "E": "Estado",
    "F": "Mi Última Puja (EUR)",
    "G": "Puja Actual/Final (EUR)",
    "H": "Reserva Alcanzada",
    "I": "Techo Dinámico (EUR)",
    "J": "Notas / Último Evento",
}


@dataclass
class WorkbookConfigSnapshot:
    """Resultado de leer la configuración de entrada desde el archivo Excel."""

    global_budget_eur: int
    lots: list[WatchLot]


@dataclass(frozen=True)
class ExcelInputSignature:
    """Firma inmutable de las entradas del usuario (B1 y columnas A:B desde fila 7)."""

    global_budget_eur: int
    lots_inputs: tuple[tuple[str, int, int], ...]  # ((url, max_bid_eur, row_idx), ...)


def extract_input_signature(file_path: str | Path) -> ExcelInputSignature | None:
    """
    Lee de forma ligera las celdas de entrada del usuario (B1 y columnas A:B desde fila 7).
    Devuelve ExcelInputSignature, o None si el archivo no existe, está bloqueado
    (PermissionError/OSError) o contiene datos temporales no parseables.
    """
    target = Path(file_path)
    if not target.exists():
        return None
    try:
        wb = load_workbook(target, data_only=True, read_only=True)
    except (PermissionError, OSError):
        return None
    except Exception:
        return None

    try:
        ws = wb.active
        if ws is None:
            return None
        raw_budget = ws["B1"].value
        budget = _parse_positive_eur(raw_budget)
        if budget is None:
            return None

        inputs: list[tuple[str, int, int]] = []
        max_row = max(ws.max_row or FIRST_DATA_ROW, FIRST_DATA_ROW - 1)
        for r in range(FIRST_DATA_ROW, max_row + 1):
            raw_url = ws.cell(row=r, column=1).value
            raw_max_bid = ws.cell(row=r, column=2).value
            if (raw_url is None or str(raw_url).strip() == "") and (
                raw_max_bid is None or str(raw_max_bid).strip() == ""
            ):
                continue
            url_str = str(raw_url).strip() if raw_url is not None else ""
            max_bid = _parse_positive_eur(raw_max_bid) or 0
            inputs.append((url_str, max_bid, r))

        return ExcelInputSignature(
            global_budget_eur=budget,
            lots_inputs=tuple(inputs),
        )
    except (PermissionError, OSError):
        return None
    except Exception:
        return None
    finally:
        try:
            wb.close()
        except Exception:
            pass


def make_input_signature_from_memory(
    global_budget_eur: int,
    lots: list[WatchLot],
) -> ExcelInputSignature:
    """Construye una firma de entrada equivalente a partir del estado en memoria."""
    inputs = tuple(
        (lot.url, lot.max_bid_eur, lot.row_index if lot.row_index is not None else 0)
        for lot in lots
    )
    return ExcelInputSignature(
        global_budget_eur=global_budget_eur,
        lots_inputs=inputs,
    )


def _parse_positive_eur(value: Any) -> int | None:
    """Convierte un valor de celda en un entero positivo en euros, o None si es inválido."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value if value > 0 else None
    if isinstance(value, float):
        if value <= 0 or not value.is_integer():
            return int(round(value)) if value > 0 else None
        return int(value)
    if isinstance(value, str):
        cleaned = (
            value.strip()
            .replace("€", "")
            .replace("EUR", "")
            .replace("eur", "")
            .replace(" ", "")
            .replace(",", ".")
        )
        if not cleaned:
            return None
        try:
            num = float(cleaned)
        except ValueError:
            return None
        if num <= 0:
            return None
        return int(round(num))
    return None


def create_excel_template(
    file_path: str | Path,
    *,
    default_budget_eur: int = 1500,
    sample_rows: list[tuple[str, int]] | None = None,
) -> Path:
    """Crea un fichero Excel (.xlsx) con el formato oficial de ChronoBid."""
    target = Path(file_path)
    target.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws: Worksheet = wb.active  # type: ignore[assignment]
    ws.title = "Subastas Catawiki"

    bold_font = Font(bold=True)
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    input_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    output_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    # Bloque de presupuesto (Filas 1 a 4)
    budget_labels = [
        (1, "Presupuesto Global (EUR)", default_budget_eur, input_fill),
        (2, "Gastado en Firme (EUR)", 0, output_fill),
        (3, "Comprometido Líder (EUR)", 0, output_fill),
        (4, "Saldo Disponible (EUR)", default_budget_eur, output_fill),
    ]
    for row_idx, label, val, fill in budget_labels:
        cell_a = ws.cell(row=row_idx, column=1, value=label)
        cell_a.font = bold_font
        cell_b = ws.cell(row=row_idx, column=2, value=val)
        cell_b.font = bold_font
        cell_b.fill = fill

    # Cabeceras de tabla en Fila 6
    for col_idx, (col_letter, header_text) in enumerate(TABLE_HEADERS.items(), start=1):
        cell = ws.cell(row=HEADER_ROW, column=col_idx, value=header_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[col_letter].width = 24 if col_letter not in {"A", "C", "J"} else 42

    if sample_rows:
        for offset, (url, max_bid) in enumerate(sample_rows):
            r = FIRST_DATA_ROW + offset
            ws.cell(row=r, column=1, value=url)
            ws.cell(row=r, column=2, value=max_bid)
            ws.cell(row=r, column=5, value=AuctionStatus.PENDING.value)
            ws.cell(row=r, column=8, value=ReserveStatus.UNKNOWN.value)
            ws.cell(row=r, column=9, value=max_bid)

    wb.save(target)
    wb.close()
    return target


def read_workbook_config(file_path: str | Path) -> WorkbookConfigSnapshot:
    """Lee el Presupuesto Global (B1) y las filas de relojes (A:B desde fila 7)."""
    target = Path(file_path)
    if not target.exists():
        raise FileNotFoundError(f"No existe el fichero Excel: {target}")

    wb = load_workbook(target, data_only=True)
    try:
        ws: Worksheet = wb.active  # type: ignore[assignment]
        raw_budget = ws["B1"].value
        global_budget = _parse_positive_eur(raw_budget)
        if global_budget is None:
            raise ValueError(
                f"El Presupuesto Global en la celda B1 ({raw_budget!r}) debe ser un número mayor que 0."
            )

        lots: list[WatchLot] = []
        max_row = max(ws.max_row, FIRST_DATA_ROW - 1)
        for row_idx in range(FIRST_DATA_ROW, max_row + 1):
            raw_url = ws.cell(row=row_idx, column=1).value
            raw_max_bid = ws.cell(row=row_idx, column=2).value

            # Ignorar filas completamente vacías en las columnas de entrada A y B
            if (raw_url is None or str(raw_url).strip() == "") and (
                raw_max_bid is None or str(raw_max_bid).strip() == ""
            ):
                continue

            url_str = str(raw_url).strip() if raw_url is not None else ""
            parsed_max_bid = _parse_positive_eur(raw_max_bid)
            lot_id = extract_lot_id(url_str, fallback_row=row_idx)

            if not is_valid_catawiki_url(url_str) or parsed_max_bid is None:
                invalid_lot = WatchLot(
                    lot_id=lot_id,
                    url=url_str,
                    max_bid_eur=parsed_max_bid if parsed_max_bid is not None else 0,
                    row_index=row_idx,
                    status=AuctionStatus.VALIDATION_ERROR,
                    notes="Fila inválida: revise URL de Catawiki y Puja Máxima (> 0 EUR)",
                )
                lots.append(invalid_lot)
            else:
                # Preservar valores previos de salida si existen en el libro
                existing_title = str(ws.cell(row=row_idx, column=3).value or "").strip()
                if any(phrase in existing_title.lower() for phrase in ERROR_TITLE_PHRASES):
                    existing_title = ""
                existing_my_bid = _parse_positive_eur(ws.cell(row=row_idx, column=6).value) or 0
                existing_cur_bid = _parse_positive_eur(ws.cell(row=row_idx, column=7).value) or 0
                valid_lot = WatchLot(
                    lot_id=lot_id,
                    url=url_str,
                    max_bid_eur=parsed_max_bid,
                    row_index=row_idx,
                    title=existing_title,
                    status=AuctionStatus.PENDING,
                    my_last_bid_eur=existing_my_bid,
                    current_bid_eur=existing_cur_bid,
                    dynamic_ceiling_eur=parsed_max_bid,
                )
                lots.append(valid_lot)

        return WorkbookConfigSnapshot(global_budget_eur=global_budget, lots=lots)
    finally:
        wb.close()


def write_workbook_state(
    file_path: str | Path,
    budget_state: GlobalBudgetState,
    lots: list[WatchLot],
) -> None:
    """
    Escribe los contadores de presupuesto (B2:B4) y las columnas de estado (C:J)
    preservando intactas las celdas de entrada del usuario (B1 y columnas A:B).
    """
    target = Path(file_path)
    if not target.exists():
        create_excel_template(target, default_budget_eur=budget_state.global_budget_eur)

    wb = load_workbook(target)
    try:
        ws: Worksheet = wb.active  # type: ignore[assignment]

        # Actualizar únicamente celdas de salida del presupuesto (B2:B4), respetando B1
        ws["B2"].value = budget_state.spent_eur
        ws["B3"].value = budget_state.committed_leader_eur
        ws["B4"].value = budget_state.available_eur

        # Construir índice de filas por URL en disco por si algún lote no tenía row_index
        url_to_row: dict[str, int] = {}
        for r in range(FIRST_DATA_ROW, ws.max_row + 1):
            cell_url = ws.cell(row=r, column=1).value
            if cell_url is not None and str(cell_url).strip():
                url_to_row[str(cell_url).strip()] = r

        next_free_row = max(ws.max_row + 1, FIRST_DATA_ROW)

        for lot in lots:
            row_idx = lot.row_index
            if row_idx is None:
                row_idx = url_to_row.get(lot.url)
            if row_idx is None:
                row_idx = next_free_row
                next_free_row += 1
                lot.row_index = row_idx

            # No sobrescribir nunca A y B si el usuario ya tiene valores en esa fila
            if ws.cell(row=row_idx, column=1).value is None and lot.url:
                ws.cell(row=row_idx, column=1, value=lot.url)
            if ws.cell(row=row_idx, column=2).value is None and lot.max_bid_eur > 0:
                ws.cell(row=row_idx, column=2, value=lot.max_bid_eur)

            closing_str = (
                lot.closing_time.strftime("%Y-%m-%d %H:%M:%S")
                if isinstance(lot.closing_time, datetime)
                else ""
            )

            ws.cell(row=row_idx, column=3, value=lot.title)
            ws.cell(row=row_idx, column=4, value=closing_str)
            ws.cell(row=row_idx, column=5, value=lot.status.value)
            ws.cell(row=row_idx, column=6, value=lot.my_last_bid_eur)
            ws.cell(row=row_idx, column=7, value=lot.current_bid_eur)
            ws.cell(row=row_idx, column=8, value=lot.reserve_status.value)
            ws.cell(row=row_idx, column=9, value=lot.effective_ceiling_eur)
            ws.cell(row=row_idx, column=10, value=lot.notes)

        # Escritura atómica mediante fichero temporal en el mismo directorio
        fd, tmp_name = tempfile.mkstemp(
            prefix=f".{target.stem}_",
            suffix=".xlsx.tmp",
            dir=str(target.parent),
        )
        os.close(fd)
        tmp_path = Path(tmp_name)
        try:
            wb.save(tmp_path)
            os.replace(tmp_path, target)
        except Exception:
            if tmp_path.exists():
                try:
                    tmp_path.unlink()
                except OSError:
                    pass
            raise
    finally:
        wb.close()
