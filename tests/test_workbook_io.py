"""Tests for the Excel generator, reader, and writer in workbook_io.py."""

from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from chronobid.excel.workbook_io import (
    FIRST_DATA_ROW,
    create_excel_template,
    extract_input_signature,
    make_input_signature_from_memory,
    read_workbook_config,
    write_workbook_state,
)
from chronobid.models.auction import (
    AuctionStatus,
    GlobalBudgetState,
    ReserveStatus,
    WatchLot,
)


def test_create_template_and_read_valid_rows(tmp_path: Path) -> None:
    excel_path = tmp_path / "relojes.xlsx"
    sample_rows = [
        ("https://www.catawiki.com/es/l/100001-omega-constellation", 600),
        ("https://www.catawiki.com/es/l/100002-seiko-presage", 350),
    ]
    create_excel_template(excel_path, default_budget_eur=1200, sample_rows=sample_rows)
    assert excel_path.exists()

    snapshot = read_workbook_config(excel_path)
    assert snapshot.global_budget_eur == 1200
    assert len(snapshot.lots) == 2

    lot1, lot2 = snapshot.lots
    assert lot1.lot_id == "lot-100001"
    assert lot1.max_bid_eur == 600
    assert lot1.row_index == FIRST_DATA_ROW
    assert lot1.status == AuctionStatus.PENDING

    assert lot2.lot_id == "lot-100002"
    assert lot2.max_bid_eur == 350
    assert lot2.row_index == FIRST_DATA_ROW + 1


def test_read_workbook_detects_invalid_rows_and_continues(tmp_path: Path) -> None:
    excel_path = tmp_path / "relojes_mixto.xlsx"
    create_excel_template(excel_path, default_budget_eur=800)

    wb = load_workbook(excel_path)
    ws = wb.active
    assert ws is not None
    # Row 7: valid
    ws.cell(row=7, column=1, value="https://www.catawiki.com/es/l/200001-longines")
    ws.cell(row=7, column=2, value=400)
    # Row 8: missing URL
    ws.cell(row=8, column=1, value="")
    ws.cell(row=8, column=2, value=250)
    # Row 9: negative max bid
    ws.cell(row=9, column=1, value="https://www.catawiki.com/es/l/200003-tissot")
    ws.cell(row=9, column=2, value=-10)
    # Row 10: non-numeric max bid
    ws.cell(row=10, column=1, value="https://www.catawiki.com/es/l/200004-hamilton")
    ws.cell(row=10, column=2, value="abc")
    # Row 11: empty (should be ignored)
    # Row 12: valid with euro symbol
    ws.cell(row=12, column=1, value="https://www.catawiki.com/es/l/200005-certina")
    ws.cell(row=12, column=2, value="275 €")
    wb.save(excel_path)
    wb.close()

    snapshot = read_workbook_config(excel_path)
    assert snapshot.global_budget_eur == 800
    assert len(snapshot.lots) == 5

    assert snapshot.lots[0].status == AuctionStatus.PENDING
    assert snapshot.lots[0].max_bid_eur == 400

    assert snapshot.lots[1].status == AuctionStatus.VALIDATION_ERROR
    assert snapshot.lots[1].row_index == 8

    assert snapshot.lots[2].status == AuctionStatus.VALIDATION_ERROR
    assert snapshot.lots[2].row_index == 9

    assert snapshot.lots[3].status == AuctionStatus.VALIDATION_ERROR
    assert snapshot.lots[3].row_index == 10

    assert snapshot.lots[4].status == AuctionStatus.PENDING
    assert snapshot.lots[4].max_bid_eur == 275
    assert snapshot.lots[4].row_index == 12


def test_read_workbook_invalid_global_budget_raises(tmp_path: Path) -> None:
    excel_path = tmp_path / "presupuesto_invalido.xlsx"
    create_excel_template(excel_path, default_budget_eur=500)
    wb = load_workbook(excel_path)
    ws = wb.active
    assert ws is not None
    ws["B1"].value = 0
    wb.save(excel_path)
    wb.close()

    with pytest.raises(ValueError, match="Presupuesto Global"):
        read_workbook_config(excel_path)


def test_write_workbook_state_updates_outputs_and_preserves_user_inputs(tmp_path: Path) -> None:
    excel_path = tmp_path / "estado.xlsx"
    sample_rows = [
        ("https://www.catawiki.com/es/l/300001-omega", 500),
    ]
    create_excel_template(excel_path, default_budget_eur=1000, sample_rows=sample_rows)

    # Simulate user manually editing B1 and B7 on disk before the agent writes
    wb = load_workbook(excel_path)
    ws = wb.active
    assert ws is not None
    ws["B1"].value = 1250
    ws["B7"].value = 550
    wb.save(excel_path)
    wb.close()

    lot = WatchLot(
        lot_id="lot-300001",
        url="https://www.catawiki.com/es/l/300001-omega",
        max_bid_eur=500,
        row_index=7,
        title="Omega Geneve Automatic 1972",
        closing_time=datetime(2026, 10, 17, 21, 15, 0),
        status=AuctionStatus.LEADER,
        my_last_bid_eur=320,
        current_bid_eur=320,
        reserve_status=ReserveStatus.MET,
        dynamic_ceiling_eur=480,
        notes="Puja líder aceptada",
    )
    budget = GlobalBudgetState(
        global_budget_eur=1000,
        spent_eur=200,
        committed_leader_eur=320,
    )

    write_workbook_state(excel_path, budget, [lot])

    wb_check = load_workbook(excel_path, data_only=True)
    ws_check = wb_check.active
    assert ws_check is not None
    # User inputs preserved intact
    assert ws_check["B1"].value == 1250
    assert ws_check["A7"].value == "https://www.catawiki.com/es/l/300001-omega"
    assert ws_check["B7"].value == 550

    # Budget outputs updated
    assert ws_check["B2"].value == 200
    assert ws_check["B3"].value == 320
    assert ws_check["B4"].value == 480

    # Columns C..J updated
    assert ws_check["C7"].value == "Omega Geneve Automatic 1972"
    assert ws_check["D7"].value == "2026-10-17 21:15:00"
    assert ws_check["E7"].value == "LIDER"
    assert ws_check["F7"].value == 320
    assert ws_check["G7"].value == 320
    assert ws_check["H7"].value == "ALCANZADA"
    assert ws_check["I7"].value == 480
    assert ws_check["J7"].value == "Puja líder aceptada"
    wb_check.close()


def test_read_workbook_config_ignores_error_titles(tmp_path: Path) -> None:
    excel_path = tmp_path / "relojes_con_error.xlsx"
    create_excel_template(
        excel_path,
        default_budget_eur=1000,
        sample_rows=[("https://www.catawiki.com/es/l/107658893-seiko-presage", 400)],
    )

    # Simulate row 7 having the error message "We can't load this page right now" in column C
    wb = load_workbook(excel_path)
    ws = wb.active
    assert ws is not None
    ws.cell(row=7, column=3, value="We can't load this page right now")
    wb.save(excel_path)
    wb.close()

    config = read_workbook_config(excel_path)
    assert len(config.lots) == 1
    lot = config.lots[0]
    assert lot.title == ""  # Should have ignored the error message


def test_extract_input_signature_detects_changes(tmp_path: Path) -> None:
    excel_path = tmp_path / "signature_test.xlsx"
    rows = [
        ("https://www.catawiki.com/es/l/1001-omega", 500),
        ("https://www.catawiki.com/es/l/1002-seiko", 300),
    ]
    create_excel_template(excel_path, default_budget_eur=1000, sample_rows=rows)

    cfg = read_workbook_config(excel_path)
    sig_disk = extract_input_signature(excel_path)
    assert sig_disk is not None
    assert sig_disk.global_budget_eur == 1000
    assert len(sig_disk.lots_inputs) == 2

    sig_mem = make_input_signature_from_memory(cfg.global_budget_eur, cfg.lots)
    assert sig_disk == sig_mem

    # Modify B1 in Excel -> changes signature
    wb = load_workbook(excel_path)
    ws = wb.active
    assert ws is not None
    ws["B1"].value = 1500
    wb.save(excel_path)
    wb.close()

    sig_new = extract_input_signature(excel_path)
    assert sig_new is not None
    assert sig_new != sig_mem
    assert sig_new.global_budget_eur == 1500

    # Add new row -> changes signature
    wb = load_workbook(excel_path)
    ws = wb.active
    assert ws is not None
    ws.cell(row=9, column=1, value="https://www.catawiki.com/es/l/1003-tissot")
    ws.cell(row=9, column=2, value=200)
    wb.save(excel_path)
    wb.close()

    sig_added = extract_input_signature(excel_path)
    assert sig_added is not None
    assert len(sig_added.lots_inputs) == 3
    assert sig_added != sig_new

    # Non-existent file returns None
    assert extract_input_signature(tmp_path / "non_existent.xlsx") is None


def test_extract_input_signature_ignores_stale_xml_dimension(tmp_path: Path) -> None:
    """
    Verify that extract_input_signature reads all actual worksheet rows even when
    a spreadsheet editor saves the .xlsx file with a stale or truncated <dimension> tag.
    """
    import re
    import zipfile

    excel_path = tmp_path / "stale_dimension.xlsx"
    rows = [
        ("https://www.catawiki.com/es/l/1001-omega", 500),
        ("https://www.catawiki.com/es/l/1002-seiko", 300),
        ("https://www.catawiki.com/es/l/1003-tissot", 250),
    ]
    create_excel_template(excel_path, default_budget_eur=1500, sample_rows=rows)

    # Rewrite xl/worksheets/sheet1.xml inside the .xlsx archive with a stale <dimension ref="A1:J7"/>
    patched_path = tmp_path / "stale_dimension_patched.xlsx"
    with zipfile.ZipFile(excel_path, "r") as zin, zipfile.ZipFile(patched_path, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml_text = data.decode("utf-8")
                xml_text = re.sub(
                    r'<dimension\s+ref="[^"]+"\s*/>',
                    '<dimension ref="A1:J7"/>',
                    xml_text,
                )
                data = xml_text.encode("utf-8")
            zout.writestr(item, data)

    sig = extract_input_signature(patched_path)
    assert sig is not None
    assert sig.global_budget_eur == 1500
    assert len(sig.lots_inputs) == 3
    assert sig.lots_inputs[0] == ("https://www.catawiki.com/es/l/1001-omega", 500, 7)
    assert sig.lots_inputs[1] == ("https://www.catawiki.com/es/l/1002-seiko", 300, 8)
    assert sig.lots_inputs[2] == ("https://www.catawiki.com/es/l/1003-tissot", 250, 9)
