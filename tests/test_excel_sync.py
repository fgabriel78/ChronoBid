"""Tests asíncronos de ExcelSyncWorker con bloqueo de archivo simulado."""

import asyncio
from pathlib import Path

from openpyxl import load_workbook

from clock_agent.excel.sync_worker import ExcelSyncWorker
from clock_agent.excel.workbook_io import create_excel_template, write_workbook_state
from clock_agent.models.auction import (
    AuctionStatus,
    GlobalBudgetState,
    ReserveStatus,
    WatchLot,
)


async def test_excel_sync_worker_persists_when_unlocked(tmp_path: Path) -> None:
    excel_path = tmp_path / "sync_ok.xlsx"
    create_excel_template(
        excel_path,
        default_budget_eur=1000,
        sample_rows=[("https://www.catawiki.com/es/l/500001-seiko", 300)],
    )

    worker = ExcelSyncWorker(excel_path, retry_interval_seconds=0.05)
    await worker.start()

    lot = WatchLot(
        lot_id="lot-500001",
        url="https://www.catawiki.com/es/l/500001-seiko",
        max_bid_eur=300,
        row_index=7,
        title="Seiko 5 Sports",
        status=AuctionStatus.LEADER,
        my_last_bid_eur=180,
        current_bid_eur=180,
        reserve_status=ReserveStatus.NO_RESERVE,
        dynamic_ceiling_eur=300,
    )
    budget = GlobalBudgetState(global_budget_eur=1000, spent_eur=0, committed_leader_eur=180)

    worker.enqueue_state(budget, [lot])
    await asyncio.sleep(0.15)
    await worker.stop()

    assert worker.is_dirty is False
    assert worker.successful_writes_count >= 1

    wb = load_workbook(excel_path, data_only=True)
    ws = wb.active
    assert ws is not None
    assert ws["B3"].value == 180
    assert ws["B4"].value == 820
    assert ws["E7"].value == "LIDER"
    assert ws["F7"].value == 180
    wb.close()


async def test_excel_sync_worker_tolerates_permission_error_and_recovers(tmp_path: Path) -> None:
    excel_path = tmp_path / "sync_locked.xlsx"
    create_excel_template(
        excel_path,
        default_budget_eur=1000,
        sample_rows=[("https://www.catawiki.com/es/l/600001-omega", 600)],
    )

    file_locked = True

    def simulated_locked_writer(
        path: Path, budget: GlobalBudgetState, lots: list[WatchLot]
    ) -> None:
        if file_locked:
            raise PermissionError(f"[Errno 13] Permission denied: '{path}'")
        write_workbook_state(path, budget, lots)

    worker = ExcelSyncWorker(
        excel_path,
        retry_interval_seconds=0.04,
        writer=simulated_locked_writer,
    )
    await worker.start()

    lot = WatchLot(
        lot_id="lot-600001",
        url="https://www.catawiki.com/es/l/600001-omega",
        max_bid_eur=600,
        row_index=7,
        title="Omega Seamaster Cosmic",
        status=AuctionStatus.LEADER,
        my_last_bid_eur=400,
        current_bid_eur=400,
    )
    # Primera actualización mientras el archivo está bloqueado
    worker.enqueue_state(
        GlobalBudgetState(global_budget_eur=1000, spent_eur=0, committed_leader_eur=400),
        [lot],
    )
    await asyncio.sleep(0.10)

    assert worker.is_dirty is True
    assert worker.lock_Encountered_count >= 1
    assert worker.successful_writes_count == 0

    # Segunda actualización en pleno bloqueo (el usuario gana el reloj por 450 €)
    lot.status = AuctionStatus.WON
    lot.my_last_bid_eur = 450
    lot.current_bid_eur = 450
    worker.enqueue_state(
        GlobalBudgetState(global_budget_eur=1000, spent_eur=450, committed_leader_eur=0),
        [lot],
    )

    # El usuario cierra el Excel -> se libera el bloqueo
    file_locked = False
    await asyncio.sleep(0.15)
    await worker.stop()

    assert worker.is_dirty is False
    assert worker.successful_writes_count >= 1
    assert worker.last_error is None

    # Verificar que en disco quedó persistido el estado más reciente (GANADO por 450 €)
    wb = load_workbook(excel_path, data_only=True)
    ws = wb.active
    assert ws is not None
    assert ws["B2"].value == 450
    assert ws["B3"].value == 0
    assert ws["B4"].value == 550
    assert ws["E7"].value == "GANADO"
    assert ws["F7"].value == 450
    assert ws["G7"].value == 450
    wb.close()
