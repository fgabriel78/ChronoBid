"""E2E tests for assault reliability: outcome attribution, event preservation and bid retry."""

from __future__ import annotations

from collections.abc import Callable
from pathlib import Path

import pytest
from openpyxl import load_workbook

from chronobid.browser.catawiki_page import (
    HumanVerificationBlockedError,
    LotPageSnapshot,
    compute_catawiki_next_min_bid,
)
from chronobid.excel.workbook_io import create_excel_template
from chronobid.models.auction import AuctionStatus, ReserveStatus, WatchLot
from chronobid.orchestrator import AuctionOrchestrator

LOT_URL = "https://www.catawiki.com/es/l/107567796-seiko-presage-srpg05j1"


class ScriptedRoom:
    """
    Deterministic lot room: the countdown decreases by `tick` seconds on each read and
    the lot closes at 0. `submit_behavior` decides how each bid attempt ends:
      - "accept": bid registered, agent becomes leader.
      - "unconfirmed": nothing changes on the page (post snapshot does not reflect the bid).
      - "error": raises an exception (missing controls, challenge, page failure...).
    A list allows a different behavior per attempt (last one repeats).
    """

    def __init__(
        self,
        *,
        current_bid: int,
        seconds: float,
        tick: float,
        submit_behavior: list[str],
        on_read: Callable[["ScriptedRoom"], None] | None = None,
        blocked_reads: Callable[["ScriptedRoom"], bool] | None = None,
    ) -> None:
        self.current_bid = current_bid
        self.seconds = seconds
        self.tick = tick
        self.submit_behavior = submit_behavior
        self.on_read = on_read
        self.blocked_reads = blocked_reads
        self.agent_is_leader = False
        self.is_closed = False
        self.bid_attempts: list[int] = []
        self.challenge_checks = 0
        self.challenge_check_seconds: list[float] = []
        self.reads = 0

    def _snap(self) -> LotPageSnapshot:
        return LotPageSnapshot(
            title="Seiko Presage SRPG05J1",
            current_bid_eur=self.current_bid,
            next_min_bid_eur=compute_catawiki_next_min_bid(self.current_bid),
            seconds_remaining=0.0 if self.is_closed else self.seconds,
            is_leader=self.agent_is_leader,
            reserve_status=ReserveStatus.NO_RESERVE,
            is_closed=self.is_closed,
        )

    async def extract_snapshot(self) -> LotPageSnapshot:
        self.reads += 1
        if not self.is_closed:
            self.seconds = max(0.0, self.seconds - self.tick)
            if self.seconds <= 0:
                self.is_closed = True
        if self.on_read is not None:
            self.on_read(self)
        if self.blocked_reads is not None and self.blocked_reads(self):
            raise HumanVerificationBlockedError("Press & Hold sin resolver")
        return self._snap()

    async def submit_minimum_bid(self, bid_amount_eur: int) -> LotPageSnapshot:
        idx = min(len(self.bid_attempts), len(self.submit_behavior) - 1)
        behavior = self.submit_behavior[idx]
        self.bid_attempts.append(bid_amount_eur)
        if behavior == "error":
            raise RuntimeError("No se encontró el botón de envío de puja")
        if behavior == "accept":
            self.current_bid = bid_amount_eur
            self.agent_is_leader = True
        return self._snap()

    async def ensure_no_challenge(self) -> bool:
        self.challenge_checks += 1
        self.challenge_check_seconds.append(self.seconds)
        return False

    async def close(self) -> None:
        return None


def _make_orchestrator(tmp_path: Path, room: ScriptedRoom, *, budget: int = 1500, max_bid: int = 350):
    excel_path = tmp_path / "assault.xlsx"
    create_excel_template(excel_path, default_budget_eur=budget, sample_rows=[(LOT_URL, max_bid)])

    async def factory(lot: WatchLot) -> ScriptedRoom:
        return room

    orch = AuctionOrchestrator(
        excel_path,
        controller_factory=factory,
        assault_threshold_seconds=300.0,
        patrol_interval_seconds=0.01,
        assault_tick_seconds=0.001,
        snipe_window_seconds=25.0,
        excel_retry_seconds=0.01,
    )
    return orch, excel_path


@pytest.mark.asyncio
async def test_unconfirmed_bid_until_close_is_generic_loss_with_event(tmp_path: Path) -> None:
    room = ScriptedRoom(current_bid=300, seconds=60.0, tick=5.0, submit_behavior=["unconfirmed"])
    orch, excel_path = _make_orchestrator(tmp_path, room)

    await orch.run_until_complete()

    lot = orch.lots[0]
    assert room.bid_attempts, "the agent must have attempted to bid inside the snipe window"
    assert lot.status == AuctionStatus.LOST_OTHER
    assert lot.my_last_bid_eur == 0
    assert "no confirmada" in lot.notes
    assert orch.ledger.committed_leader_eur == 0

    wb = load_workbook(excel_path, data_only=True)
    ws = wb.active
    assert ws is not None
    assert ws["E7"].value == "PERDIDO"
    assert "Último evento: Puja de" in (ws["J7"].value or "")
    wb.close()


@pytest.mark.asyncio
async def test_unconfirmed_bid_is_retried_on_next_tick(tmp_path: Path) -> None:
    room = ScriptedRoom(
        current_bid=300, seconds=60.0, tick=5.0, submit_behavior=["unconfirmed", "accept"]
    )
    released_between: list[int] = []

    orch, _ = _make_orchestrator(tmp_path, room)
    original_submit = room.submit_minimum_bid

    async def tracking_submit(amount: int) -> LotPageSnapshot:
        released_between.append(orch.ledger.committed_leader_eur)
        return await original_submit(amount)

    room.submit_minimum_bid = tracking_submit  # type: ignore[method-assign]

    await orch.run_until_complete()

    assert len(room.bid_attempts) >= 2
    assert room.bid_attempts[0] == room.bid_attempts[1] == 320
    # Before the second attempt only that attempt's own reservation is committed
    assert released_between[1] == 320
    assert orch.lots[0].status == AuctionStatus.WON


@pytest.mark.asyncio
async def test_submission_error_is_retried_and_recorded(tmp_path: Path) -> None:
    room = ScriptedRoom(current_bid=300, seconds=60.0, tick=5.0, submit_behavior=["error", "accept"])
    orch, _ = _make_orchestrator(tmp_path, room)

    await orch.run_until_complete()

    lot = orch.lots[0]
    assert len(room.bid_attempts) >= 2
    assert lot.status == AuctionStatus.WON
    assert orch.ledger.spent_eur == 320


@pytest.mark.asyncio
async def test_submission_error_until_close_keeps_error_event(tmp_path: Path) -> None:
    room = ScriptedRoom(current_bid=300, seconds=60.0, tick=5.0, submit_behavior=["error"])
    orch, _ = _make_orchestrator(tmp_path, room)

    await orch.run_until_complete()

    lot = orch.lots[0]
    assert lot.status == AuctionStatus.LOST_OTHER
    assert "Error de envío de puja" in lot.notes
    assert orch.ledger.committed_leader_eur == 0


@pytest.mark.asyncio
async def test_patrol_verification_block_is_reported_and_cleared(tmp_path: Path) -> None:
    room = ScriptedRoom(
        current_bid=200,
        seconds=7200.0,
        tick=0.0,
        submit_behavior=["accept"],
        blocked_reads=lambda r: r.reads == 1,
    )
    orch, excel_path = _make_orchestrator(tmp_path, room)
    lot = orch.lots[0]

    await orch.run_patrol_step()
    assert lot.status == AuctionStatus.BLOCKED_VERIFICATION
    assert lot.status.is_active
    assert "verificación humana" in lot.notes

    await orch.sync_worker.flush_once()
    wb = load_workbook(excel_path, data_only=True)
    ws = wb.active
    assert ws is not None
    assert ws["E7"].value == "BLOQUEADO: Verificación Humana"
    assert "verificación humana" in (ws["J7"].value or "")
    wb.close()

    await orch.run_patrol_step(force=True)
    assert lot.status == AuctionStatus.PATROL
    assert lot.current_bid_eur == 200


@pytest.mark.asyncio
async def test_assault_blocked_until_close_ends_as_generic_loss(tmp_path: Path) -> None:
    room = ScriptedRoom(
        current_bid=300,
        seconds=60.0,
        tick=5.0,
        submit_behavior=["accept"],
        blocked_reads=lambda r: 0.0 < r.seconds <= 30.0,
    )
    orch, _ = _make_orchestrator(tmp_path, room)

    await orch.run_until_complete()

    lot = orch.lots[0]
    assert room.bid_attempts == []
    assert lot.status == AuctionStatus.LOST_OTHER
    assert "Último evento: Bloqueado por verificación humana" in lot.notes


@pytest.mark.asyncio
async def test_pre_snipe_check_runs_once_per_crossing_and_rearms_after_extension(
    tmp_path: Path,
) -> None:
    extended = False

    def rival_extends_once(room: ScriptedRoom) -> None:
        nonlocal extended
        if room.bid_attempts and not extended:
            extended = True
            room.seconds = 90.0
            room.current_bid = compute_catawiki_next_min_bid(room.current_bid)
            room.agent_is_leader = False

    room = ScriptedRoom(
        current_bid=300,
        seconds=60.0,
        tick=5.0,
        submit_behavior=["accept"],
        on_read=rival_extends_once,
    )
    orch, _ = _make_orchestrator(tmp_path, room, max_bid=500)

    await orch.run_until_complete()

    assert extended
    assert room.challenge_checks == 2
    for secs in room.challenge_check_seconds:
        assert 25.0 < secs <= 35.0
    assert orch.lots[0].status == AuctionStatus.WON


@pytest.mark.asyncio
async def test_lost_in_assault_emits_full_diagnostic_trace(tmp_path: Path, caplog) -> None:
    import logging

    caplog.set_level(logging.DEBUG, logger="chronobid")
    room = ScriptedRoom(current_bid=300, seconds=60.0, tick=5.0, submit_behavior=["unconfirmed"])
    orch, _ = _make_orchestrator(tmp_path, room)

    await orch.run_until_complete()

    messages = [r.getMessage() for r in caplog.records if r.name.startswith("chronobid")]
    joined = "\n".join(messages)
    assert "SESSION_START" in joined and "SESSION_LOT row=7" in joined
    assert "ASSAULT_ENTRY lot=" in joined
    assert "ASSAULT_EXIT lot=" in joined
    decisions = [m for m in messages if m.startswith("DECISION")]
    # Every tick inside the snipe window is traced
    assert sum(1 for m in decisions if "window=True" in m) >= 3
    assert any("BID_ATTEMPT" in m and "amount=320" in m for m in messages)
    assert any("BID_RESULT" in m and "confirmed=False" in m for m in messages)
    finals = [m for m in messages if m.startswith("FINAL_OUTCOME")]
    assert len(finals) == 1
    assert "status=PERDIDO " in finals[0]
    assert "no confirmada" in finals[0]
